"""Local-only NetMaxx trial. No paid API calls or automatic LinkedIn scraping."""
import base64,csv,io,json,re,sys
from pathlib import Path
from http.server import ThreadingHTTPServer,BaseHTTPRequestHandler
from urllib.parse import urlsplit
ROOT=Path(__file__).resolve().parents[1]
THEMES={
 'Product & design':['product','design','ux','roadmap','user discovery','users','prototyp'],
 'Deployment & delivery':['deploy','delivery','integration','integrat','workflow','field','stakeholder'],
 'AI & data':['artificial intelligence','ai','ml','machine learning','data','analytics','cnn','model'],
 'Industrial systems':['industrial','manufactur','rail','transport','robot','autonomy','systems','energy'],
 'Strategy & research':['strategy','strategic','research','market','business','insight','gtm'],
 'Community & leadership':['mentor','workshop','community','leadership','cross-functional','team','facilitat']}

def themes(text):
 text=text.lower()
 return [name for name,words in THEMES.items() if any(re.search(r'\b'+re.escape(w)+(r'\b' if len(w)<=3 else ''),text) for w in words)]

def slug(url):
 try:
  p=urlsplit(url if '://' in url else 'https://'+url)
  return p.path.rstrip('/').lower() if (p.hostname or '').endswith('.linkedin.com') or p.hostname=='linkedin.com' else None
 except ValueError:return None

def sample():
 report=json.loads((ROOT/'private/netmaxx_public_review.json').read_text())
 profile=report['icp']['requester']['background']+' Skills: '+', '.join(report['icp']['requester']['skills'])
 return {'profile':profile,'objective':'Meet people doing work similar to my experience.','attendees':json.loads((ROOT/'private/netmaxx_attendees.json').read_text()),'review':report}

def parse_file(payload):
 raw=base64.b64decode(payload['data'],validate=True)
 if len(raw)>5_000_000:raise ValueError('Please choose a file smaller than 5 MB.')
 name=payload['name'].lower()
 if payload['kind']=='resume':
  if name.endswith('.pdf'):
   from pypdf import PdfReader
   text='\n'.join(p.extract_text() or '' for p in PdfReader(io.BytesIO(raw)).pages[:15])
  elif name.endswith('.txt'):text=raw.decode('utf-8-sig')
  else:raise ValueError('Use a PDF or plain text resume.')
  if len(text.strip())<40:raise ValueError('No readable resume text found. Please paste your professional background instead.')
  return {'text':text[:30000]}
 if name.endswith('.xlsx'):
  from openpyxl import load_workbook
  book=load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
  sheet=book.worksheets[0]
  if sheet.max_row>501 or sheet.max_column>30:raise ValueError('For this trial, use at most 500 rows and 30 columns.')
  values=list(sheet.iter_rows(values_only=True));book.close()
  headers=values[0];rows=[dict(zip(headers,row)) for row in values[1:] if any(x is not None for x in row)]
 elif name.endswith('.csv'):rows=list(csv.DictReader(io.StringIO(raw.decode('utf-8-sig'))))
 else:raise ValueError('Use an Excel .xlsx or CSV attendee file.')
 aliases={'name':'Attendee Name','attendee name':'Attendee Name','job title':'Job Title','current company / school':'Current Company','current company':'Current Company','company':'Current Company','linkedin':'LinkedIn URL','linkedin url':'LinkedIn URL'}
 attendees=[]
 for i,row in enumerate(rows):
  mapped={aliases[str(k).strip().lower()]:str(v).strip() if v is not None else '' for k,v in row.items() if str(k).strip().lower() in aliases}
  if not mapped.get('Attendee Name'):raise ValueError(f'Row {i+2} needs an attendee name. Use Name, Job Title, Current Company, LinkedIn URL columns.')
  attendees.append(mapped)
 if not 1<=len(attendees)<=30:raise ValueError('Choose 1–30 attendees for this trial.')
 return {'attendees':attendees}

MODE_LABELS={'peers':'Professional Peers','hiring':'Hiring Connections','social':'Social Connection'}

def interpret(objective, selected='auto'):
 text=objective.lower()
 # Remove simple explicit negatives before recognizing positive intent.
 text=re.sub(r"\b(?:not|no|nothing for|not looking for)\s+(?:interested in\s+)?(?:dating|dates|hiring|jobs?|friendship)\b", '', text)
 hiring=bool(re.search(r'\b(hir(?:e|ing)|jobs?|recruit(?:er|ers|ing)?|employ(?:er|ment)|career opportunities)\b',text))
 social=bool(re.search(r'\b(dating|dates?|friends?|friendship|social|fun|jokes?|memes?|humou?r)\b',text))
 peers=bool(re.search(r'\b(peers?|professional|experience|similar|connect|work|background)\b',text))
 if selected not in {'auto',*MODE_LABELS}:raise ValueError('Choose Hiring, Professional Peers or Social Connection.')
 if selected=='auto':
  if hiring and social:raise ValueError('Your request includes hiring and social goals. Choose one mode to compare first.')
  if not (hiring or social or peers) and not re.search(r'\b(women|female|men|male|nonbinary)\b',text):raise ValueError('Please choose a mode, or describe hiring, professional-peer or social connections.')
 mode=selected if selected!='auto' else 'hiring' if hiring else 'social' if social else 'peers'
 gender=bool(re.search(r'\b(women|woman|female|men|man|male|nonbinary|non-binary|opposite gender)\b',text))
 return mode,gender

def rank(payload):
 background=payload.get('profile','').strip();objective=payload.get('objective','').strip()
 mode,gender_filter=interpret(objective,payload.get('mode','auto'))
 if mode!='social' and len(background)<40:raise ValueError('Add your resume or a few sentences about your work for professional matching.')
 candidates=payload.get('attendees',[])
 if not isinstance(candidates,list) or not 1<=len(candidates)<=30:raise ValueError('Add 1–30 attendees first.')
 lookup={slug(r['LinkedIn URL']):r for r in sample()['review']['attendees']}
 target=themes(background);desired=themes(objective)
 interests=(payload.get('interests','')+' '+objective).lower()
 results=[];seen=set()
 for person in candidates:
  key=slug(person.get('LinkedIn URL',''))
  if key and key in seen:continue
  seen.add(key);known=lookup.get(key) if key else None
  result={**person,'score':None,'matches':[],'evidence':[],'gap':'No researched public evidence is available for this profile yet.','status':'Needs evidence','reason':'Cannot assess this person without supporting evidence.','eligible':False,'score_label':'Relevance'}
  if known:
   evidence=known['evidence'];result.update(evidence=evidence,gap=known['data_gaps'][0])
   if known['score'] is None:result.update(status=known['status'],reason=known['reason'])
   elif person.get('Current Company','').strip().casefold()!=known['Current Company'].strip().casefold():result.update(status='Check company',gap='Uploaded company differs from the researched record. Confirm identity before ranking.')
   else:
    available=themes(' '.join(x['summary'] for x in evidence if x['id'].startswith('e')))
    overlap=[x for x in target if x in available]
    weighted={x:2 if x in desired else 1 for x in target}
    fit=sum(weighted[x] for x in overlap)/max(1,sum(weighted.values()))
    signals=known.get('intent_signals',{})
    if mode=='peers':
     if overlap:result.update(score=round(10*fit,1),eligible=True,status='Provisional professional overlap',matches=overlap,reason='Evidence supports these shared work themes: '+', '.join(overlap)+'.',missing=[x for x in target if x not in available])
     else:result.update(status='No evidenced work overlap',reason='No shared professional themes were established from the available evidence.')
    elif mode=='hiring':
     hiring=signals.get('hiring')
     if hiring and overlap:
      points=6 if hiring['kind']=='recruiter' else 3
      result.update(score=round(points+3*fit,1),eligible=True,status='Recruiting connection · opening unconfirmed' if points==6 else 'Founder connection · hiring unconfirmed',matches=[hiring['kind'].title(),*overlap],reason=hiring['summary'],gap='Hiring evidence may be historical. Current vacancies, role fit, budget and willingness to refer or hire you require confirmation. '+known['data_gaps'][0])
     else:result.update(status='Hiring influence not established',reason='Professional similarity or seniority alone does not show that this person can hire you.',gap='No supported hiring or founder signal with relevant work overlap was found.')
    else:
     social=signals.get('social',[])
     if social:
      matching=[x for x in social if any(re.search(r'\b'+re.escape(w)+r'\b',interests) for w in x['keywords'])]
      chosen=matching or social
      explicit=bool(payload.get('interests','').strip()) or bool(re.search(r'\b(jazz|music|humou?r|jokes?|memes?|funny|hiking|sports|art|food|travel)\b',objective.lower()))
      if explicit and not matching:result.update(status='No evidenced interest overlap',reason='Public social cues exist, but do not match the interests you specified.',gap='Unknown interests are not negative personality judgments.')
      else:
       ids={i for x in chosen for i in x['evidence_ids']}
       result.update(score=min(10,max(x['strength'] for x in chosen)+(3 if matching else 0)),eligible=True,status='Public conversation cue · not compatibility',matches=[x['topic'] for x in chosen],reason=' '.join(x['summary'] for x in chosen),evidence=[x for x in evidence if x['id'] in ids],gap='Your social preferences are not supplied; this orders visible conversation cues only.' if not matching else 'Shared interest cues do not establish chemistry, current interests or romantic intent.',score_label='Conversation evidence')
     else:result.update(status='No public social cues found',reason='Professional achievements are not used as a substitute for hobbies, humor or social interests.',gap='No attributable social-interest evidence was collected. No X account was confidently linked.')
  if gender_filter:result.update(score=None,eligible=False,matches=[],status='Gender filter cannot be established',reason='Self-described gender is unavailable in this evidence set. No gender was inferred from names, photos or pronouns.',gap='No confirmed matches for this filter; add voluntarily supplied information to apply it.')
  results.append(result)
 results.sort(key=lambda r:(not r['eligible'],-(r['score'] or 0),r['Attendee Name'].lower()))
 return {'results':results,'mode':mode,'interpretation':MODE_LABELS[mode]+(' · gender constraint requested; no confirmed data' if gender_filter else ''),'themes':target,'top_candidate_urls':[r.get('LinkedIn URL') for r in results if r['eligible']][:3], 'method':'Intent-specific evidence rules',
 'note':('Social mode uses public conversation cues, not achievements or romantic availability. Without your interests this is not personalized compatibility. ' if mode=='social' else 'Different rules apply to hiring connections and professional peers. ')+ 'Scores are explainable heuristics, not confidence percentages. Public evidence is a saved research snapshot; dates and current roles may be incomplete. No paid API calls.'}

class Handler(BaseHTTPRequestHandler):
 def log_message(self,*args):pass
 def respond(self,status,data,ctype='application/json'):
  body=json.dumps(data).encode() if ctype=='application/json' else data
  self.send_response(status);self.send_header('Content-Type',ctype);self.send_header('Content-Length',str(len(body)));self.send_header('Cache-Control','no-store');self.send_header('X-Content-Type-Options','nosniff');self.end_headers();self.wfile.write(body)
 def valid_host(self):return self.headers.get('Host') in {'127.0.0.1:8765','localhost:8765'}
 def do_GET(self):
  if not self.valid_host():return self.respond(403,{'error':'Local access only'})
  if self.path=='/api/sample':return self.respond(200,sample())
  if self.path in ('/','/index.html'):return self.respond(200,(ROOT/'app/index.html').read_bytes(),'text/html; charset=utf-8')
  self.respond(404,{'error':'Not found'})
 def do_POST(self):
  if not self.valid_host() or self.headers.get('Origin') not in {'http://127.0.0.1:8765','http://localhost:8765'}:return self.respond(403,{'error':'Local access only'})
  try:
   n=int(self.headers.get('Content-Length','0'))
   if not 0<n<7_000_000:raise ValueError('File or request is too large.')
   data=json.loads(self.rfile.read(n))
   if self.path=='/api/parse':result=parse_file(data)
   elif self.path=='/api/rank':result=rank(data)
   else:return self.respond(404,{'error':'Not found'})
   self.respond(200,result)
  except Exception as exc:
   self.respond(400,{'error':str(exc) if isinstance(exc,ValueError) else 'Could not read this file. Check its format and try again.'})

if __name__=='__main__':
 print('NetMaxx running at http://127.0.0.1:8765',flush=True)
 ThreadingHTTPServer(('127.0.0.1',8765),Handler).serve_forever()
